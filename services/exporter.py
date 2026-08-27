"""导出 Excel。"""
from __future__ import annotations

from datetime import datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import func

from extensions import db
from models import ScoreRecord, Student
from services.categories import CATEGORY_MAP


HEADER_FILL = PatternFill("solid", fgColor="7C3AED")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def export_records_xlsx(semester: str) -> tuple[BytesIO, str]:
    """导出某学期所有有效加分记录 + 一张汇总表。"""
    wb = Workbook()

    # Sheet 1: 明细
    ws = wb.active
    ws.title = "加分明细"
    headers = ["学号", "姓名", "加分类别", "活动名称", "分值", "是否公示", "学期", "录入时间"]
    ws.append(headers)
    for c in ws[1]:
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center")

    rows = (db.session.query(ScoreRecord, Student)
            .join(Student, Student.id == ScoreRecord.student_id)
            .filter(ScoreRecord.semester == semester,
                    ScoreRecord.is_revoked.is_(False))
            .order_by(Student.student_no.asc(), ScoreRecord.created_at.asc())
            .all())
    for rec, st in rows:
        ws.append([
            st.student_no, st.name, CATEGORY_MAP.get(rec.category, rec.category),
            rec.activity_name, float(rec.points), "是" if rec.is_public else "否", rec.semester,
            rec.created_at.strftime("%Y-%m-%d %H:%M") if rec.created_at else "",
        ])
    for i, w in enumerate([16, 14, 24, 32, 10, 12, 14, 18], start=1):
        ws.column_dimensions[chr(64 + i)].width = w

    # Sheet 2: 汇总（按学生）
    ws2 = wb.create_sheet("学期汇总")
    ws2.append(["学号", "姓名", "总分", "记录数"])
    for c in ws2[1]:
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center")

    summary = (db.session.query(
                    Student.student_no, Student.name,
                    func.coalesce(func.sum(ScoreRecord.points), 0).label("total"),
                    func.count(ScoreRecord.id).label("cnt"))
               .outerjoin(ScoreRecord,
                          (ScoreRecord.student_id == Student.id)
                          & (ScoreRecord.semester == semester)
                          & (ScoreRecord.is_revoked.is_(False)))
               .group_by(Student.id)
               .order_by(func.coalesce(func.sum(ScoreRecord.points), 0).desc(),
                         Student.student_no.asc())
               .all())
    for no, name, total, cnt in summary:
        ws2.append([no, name, float(total or 0), int(cnt or 0)])
    for i, w in enumerate([16, 14, 10, 10], start=1):
        ws2.column_dimensions[chr(64 + i)].width = w

    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    filename = f"加分汇总-{semester}-{datetime.now().strftime('%Y%m%d%H%M')}.xlsx"
    return bio, filename


def export_roster_xlsx(semester: str) -> tuple[BytesIO, str]:
    """导出当前学生名单，作为主页上可直接下载的名单文件。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "学生名单"
    ws.append(["学号", "姓名", "学期"])
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    for student in Student.query.order_by(Student.student_no.asc()).all():
        ws.append([student.student_no, student.name, semester])
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 16
    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio, f"学生名单-{semester}.xlsx"
